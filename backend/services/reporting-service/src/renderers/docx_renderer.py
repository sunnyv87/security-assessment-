"""DOCX renderer using python-docx for Word export."""

from __future__ import annotations

import io
import re

import structlog
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.models.domain import (
    ReportDataBundle,
    ReportType,
    FindingData,
    ReportStatistics,
)

logger = structlog.get_logger()

SEVERITY_COLORS = {
    "critical": RGBColor(0xDC, 0x26, 0x26),
    "high": RGBColor(0xEA, 0x58, 0x0C),
    "medium": RGBColor(0xCA, 0x8A, 0x04),
    "low": RGBColor(0x25, 0x63, 0xEB),
    "info": RGBColor(0x6B, 0x72, 0x80),
}


class DOCXRenderer:
    """Renders a ReportDataBundle into a DOCX (Word) document."""

    def render(
        self,
        bundle: ReportDataBundle,
        report_type: ReportType,
        *,
        report_title: str,
        generated_date: str,
        report_version: int = 1,
        is_draft: bool = False,
    ) -> bytes:
        """Build a DOCX document from the report data bundle.

        Returns:
            DOCX file as bytes.
        """
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        # ── Cover page ──
        self._add_cover_page(doc, bundle, report_title, generated_date, report_version, is_draft)

        # ── Executive Summary ──
        self._add_executive_summary(doc, bundle)

        # ── Scope & Methodology ──
        self._add_scope_methodology(doc, bundle)

        # ── Findings Summary ──
        self._add_findings_summary(doc, bundle.statistics)

        # ── Detailed Findings ──
        if report_type in (ReportType.FULL_TECHNICAL, ReportType.EXECUTIVE_SUMMARY):
            self._add_detailed_findings(doc, bundle.findings)

        # ── Remediation Roadmap ──
        if report_type in (ReportType.REMEDIATION_ROADMAP, ReportType.FULL_TECHNICAL):
            self._add_remediation_roadmap(doc, bundle)

        # ── Compliance Mapping ──
        if report_type in (ReportType.COMPLIANCE, ReportType.FULL_TECHNICAL):
            self._add_compliance_mapping(doc, bundle)

        # ── Disclaimer ──
        self._add_disclaimer(doc, is_draft)

        buffer = io.BytesIO()
        doc.save(buffer)
        result = buffer.getvalue()
        logger.info("docx_render_complete", size=len(result))
        return result

    # ── Section builders ──

    def _add_cover_page(
        self,
        doc: Document,
        bundle: ReportDataBundle,
        title: str,
        date: str,
        version: int,
        is_draft: bool,
    ) -> None:
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if is_draft:
            draft_p = doc.add_paragraph()
            draft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = draft_p.add_run("DRAFT — NOT FOR DISTRIBUTION")
            run.bold = True
            run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            run.font.size = Pt(14)

        meta = [
            ("Customer", bundle.engagement.customer_name),
            ("Engagement", bundle.engagement.name),
            ("Date", date),
            ("Version", str(version)),
            ("Lead Analyst", bundle.engagement.lead_analyst),
        ]
        table = doc.add_table(rows=len(meta), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, value) in enumerate(meta):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value
            table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("CONFIDENTIAL")
        run.bold = True
        run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
        doc.add_page_break()

    def _add_executive_summary(self, doc: Document, bundle: ReportDataBundle) -> None:
        doc.add_heading("1. Executive Summary", level=1)
        stats = bundle.statistics
        doc.add_paragraph(
            f"This assessment identified {stats.total_findings} findings across "
            f"{len(stats.by_asset)} assets. "
            f"Critical: {stats.by_severity.get('critical', 0)}, "
            f"High: {stats.by_severity.get('high', 0)}, "
            f"Medium: {stats.by_severity.get('medium', 0)}, "
            f"Low: {stats.by_severity.get('low', 0)}, "
            f"Info: {stats.by_severity.get('info', 0)}."
        )

        # Key stats table
        table = doc.add_table(rows=4, cols=2, style="Light Grid Accent 1")
        rows = [
            ("Total Findings", str(stats.total_findings)),
            ("Validation Coverage", f"{stats.validation_coverage:.0%}"),
            ("False Positive Rate", f"{stats.false_positive_rate:.0%}"),
            ("Remediation Coverage", f"{stats.remediation_coverage:.0%}"),
        ]
        for i, (label, value) in enumerate(rows):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value

        # Critical finding highlights
        critical = [f for f in bundle.findings if f.severity == "critical"]
        if critical:
            doc.add_heading("Critical Findings Requiring Immediate Attention", level=2)
            for finding in critical:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"{finding.title}")
                run.bold = True
                p.add_run(f" — {finding.asset} (CVSS {finding.cvss_score})")
        doc.add_page_break()

    def _add_scope_methodology(self, doc: Document, bundle: ReportDataBundle) -> None:
        eng = bundle.engagement
        doc.add_heading("2. Scope & Methodology", level=1)

        doc.add_heading("2.1 Scope Definition", level=2)
        if eng.scope_description:
            doc.add_paragraph(eng.scope_description)

        if eng.in_scope_assets:
            doc.add_heading("In-Scope Assets", level=3)
            for asset in eng.in_scope_assets:
                doc.add_paragraph(asset, style="List Bullet")

        if eng.exclusions:
            doc.add_heading("Out of Scope", level=3)
            for excl in eng.exclusions:
                doc.add_paragraph(excl, style="List Bullet")

        doc.add_heading("2.2 Testing Methodology", level=2)
        doc.add_paragraph(
            f"Testing was conducted in accordance with the {eng.methodology} methodology. "
            "The assessment included both automated scanning and manual validation "
            "by certified security analysts."
        )

        doc.add_heading("2.3 Testing Timeline", level=2)
        table = doc.add_table(rows=3, cols=2, style="Light Grid Accent 1")
        for i, (label, value) in enumerate([
            ("Start Date", eng.start_date),
            ("End Date", eng.end_date),
            ("Duration", f"{eng.duration_days} days"),
        ]):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value

        if eng.tools_used:
            doc.add_heading("2.4 Tools Used", level=2)
            for tool in eng.tools_used:
                doc.add_paragraph(tool, style="List Bullet")
        doc.add_page_break()

    def _add_findings_summary(self, doc: Document, stats: ReportStatistics) -> None:
        doc.add_heading("3. Findings Summary", level=1)

        doc.add_heading("3.1 Severity Distribution", level=2)
        table = doc.add_table(rows=6, cols=2, style="Light Grid Accent 1")
        table.rows[0].cells[0].text = "Severity"
        table.rows[0].cells[1].text = "Count"
        for cell in table.rows[0].cells:
            cell.paragraphs[0].runs[0].bold = True
        for i, sev in enumerate(["critical", "high", "medium", "low", "info"], 1):
            count = stats.by_severity.get(sev, 0)
            table.rows[i].cells[0].text = sev.upper()
            table.rows[i].cells[1].text = str(count)
            color = SEVERITY_COLORS.get(sev)
            if color:
                table.rows[i].cells[0].paragraphs[0].runs[0].font.color.rgb = color

        if stats.by_category:
            doc.add_heading("3.2 Findings by Category", level=2)
            table = doc.add_table(rows=len(stats.by_category) + 1, cols=2, style="Light Grid Accent 1")
            table.rows[0].cells[0].text = "Category"
            table.rows[0].cells[1].text = "Count"
            for cell in table.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            for i, (cat, count) in enumerate(sorted(stats.by_category.items(), key=lambda x: -x[1]), 1):
                table.rows[i].cells[0].text = cat
                table.rows[i].cells[1].text = str(count)
        doc.add_page_break()

    def _add_detailed_findings(self, doc: Document, findings: list[FindingData]) -> None:
        doc.add_heading("4. Detailed Findings", level=1)

        severity_order = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
        sorted_findings = sorted(findings, key=lambda f: severity_order.get(f.severity, 4))

        for idx, finding in enumerate(sorted_findings, 1):
            doc.add_heading(f"4.{idx} {finding.title}", level=2)

            # Metadata table
            meta_rows = [
                ("Severity", finding.severity.upper()),
                ("CVSS Score", str(finding.cvss_score)),
                ("CWE", f"{finding.cwe_id} — {finding.cwe_name}" if finding.cwe_id else "N/A"),
                ("Asset", finding.asset),
                ("Endpoint", finding.endpoint or "N/A"),
                ("Validation", finding.validation_verdict or "Pending"),
            ]
            table = doc.add_table(rows=len(meta_rows), cols=2, style="Light Grid Accent 1")
            for i, (label, value) in enumerate(meta_rows):
                table.rows[i].cells[0].text = label
                table.rows[i].cells[1].text = value
                table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

            # Description
            doc.add_heading("Description", level=3)
            doc.add_paragraph(finding.description)

            # Business Impact
            if finding.business_impact:
                doc.add_heading("Business Impact", level=3)
                doc.add_paragraph(finding.business_impact)

            # Evidence
            if finding.evidence.request or finding.evidence.response:
                doc.add_heading("Evidence", level=3)
                if finding.evidence.request:
                    req = finding.evidence.request
                    doc.add_paragraph(f"Request: {req.method} {req.url}", style="No Spacing")
                    if req.body:
                        doc.add_paragraph(req.body, style="No Spacing")
                if finding.evidence.response:
                    resp = finding.evidence.response
                    doc.add_paragraph(f"Response: {resp.status_code}", style="No Spacing")
                    if resp.body:
                        body_preview = resp.body[:500]
                        doc.add_paragraph(body_preview, style="No Spacing")

            # Reproduction steps
            if finding.reproduction_steps:
                doc.add_heading("Reproduction Steps", level=3)
                for step in finding.reproduction_steps:
                    doc.add_paragraph(step, style="List Number")

            # Remediation
            if finding.remediation.summary:
                doc.add_heading("Remediation", level=3)
                doc.add_paragraph(finding.remediation.summary)
                if finding.remediation.steps:
                    for step in finding.remediation.steps:
                        doc.add_paragraph(step, style="List Number")
                for ex in finding.remediation.code_examples:
                    if ex.before:
                        doc.add_paragraph(f"Vulnerable ({ex.language}):", style="No Spacing")
                        doc.add_paragraph(ex.before, style="No Spacing")
                    if ex.after:
                        doc.add_paragraph(f"Secure ({ex.language}):", style="No Spacing")
                        doc.add_paragraph(ex.after, style="No Spacing")

    def _add_remediation_roadmap(self, doc: Document, bundle: ReportDataBundle) -> None:
        plan = bundle.remediation_plan
        if not plan.phases:
            return

        doc.add_heading("Remediation Roadmap", level=1)
        doc.add_paragraph(f"Total Estimated Effort: {plan.total_effort_estimate}")

        for phase in plan.phases:
            doc.add_heading(phase.name, level=2)
            doc.add_paragraph(f"Timeline: {phase.timeline}")
            doc.add_paragraph(f"Effort Estimate: {phase.effort_estimate}")
            doc.add_paragraph(f"Findings in Phase: {len(phase.finding_ids)}")
            for fid in phase.finding_ids:
                matching = [f for f in bundle.findings if f.id == fid]
                if matching:
                    doc.add_paragraph(f"{matching[0].title} ({matching[0].severity.upper()})", style="List Bullet")

        if plan.retest_recommendation:
            doc.add_heading("Retest Recommendation", level=2)
            doc.add_paragraph(plan.retest_recommendation)
        doc.add_page_break()

    def _add_compliance_mapping(self, doc: Document, bundle: ReportDataBundle) -> None:
        comp = bundle.compliance
        if not comp.frameworks:
            return

        doc.add_heading("Compliance Mapping", level=1)

        # Framework scores
        if comp.scores_by_framework:
            doc.add_heading("Framework Compliance Scores", level=2)
            table = doc.add_table(rows=len(comp.scores_by_framework) + 1, cols=2, style="Light Grid Accent 1")
            table.rows[0].cells[0].text = "Framework"
            table.rows[0].cells[1].text = "Score"
            for cell in table.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            for i, (fw, score) in enumerate(comp.scores_by_framework.items(), 1):
                table.rows[i].cells[0].text = fw
                table.rows[i].cells[1].text = f"{score:.0%}"

        # Control mappings
        if comp.control_mappings:
            doc.add_heading("Control Mapping Details", level=2)
            cols = ["Control ID", "Control", "Framework", "Status", "Findings"]
            table = doc.add_table(rows=len(comp.control_mappings) + 1, cols=len(cols), style="Light Grid Accent 1")
            for j, col in enumerate(cols):
                table.rows[0].cells[j].text = col
                table.rows[0].cells[j].paragraphs[0].runs[0].bold = True
            for i, mapping in enumerate(comp.control_mappings, 1):
                table.rows[i].cells[0].text = mapping.control_id
                table.rows[i].cells[1].text = mapping.control_name
                table.rows[i].cells[2].text = mapping.framework
                table.rows[i].cells[3].text = mapping.status.upper()
                table.rows[i].cells[4].text = str(len(mapping.finding_ids))

        # Gaps
        if comp.gaps:
            doc.add_heading("Gap Analysis", level=2)
            for gap in comp.gaps:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"{gap.control_id}: {gap.control_name}")
                run.bold = True
                p.add_run(f" ({gap.framework}) — {gap.remediation_required}")
        doc.add_page_break()

    def _add_disclaimer(self, doc: Document, is_draft: bool) -> None:
        doc.add_heading("Disclaimer", level=1)
        if is_draft:
            p = doc.add_paragraph()
            run = p.add_run(
                "AI-GENERATED DRAFT: This report was generated with AI assistance "
                "and has NOT been reviewed by a human analyst. It must be reviewed "
                "and approved before distribution."
            )
            run.bold = True
            run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

        doc.add_paragraph(
            "This report is confidential and intended solely for the named recipient(s). "
            "The findings are based on point-in-time testing and do not guarantee the "
            "absence of other vulnerabilities. Unauthorized distribution is prohibited."
        )
